from flask import Flask, render_template, request, redirect, url_for
from openpyxl import Workbook, load_workbook
import os

app = Flask(__name__)

DATA_FILE = 'data.xlsx'


def init_workbook():
    if not os.path.exists(DATA_FILE):
        wb = Workbook()
        ws_products = wb.active
        ws_products.title = 'Products'
        ws_products.append(['ID', 'Name', 'Price'])
        ws_packs = wb.create_sheet('Packs')
        ws_packs.append(['ID', 'Name', 'Items', 'Price'])
        wb.save(DATA_FILE)


def read_products():
    wb = load_workbook(DATA_FILE)
    ws = wb['Products']
    products = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            products.append({'id': row[0], 'name': row[1], 'price': row[2]})
    wb.close()
    return products


def write_products(products):
    wb = load_workbook(DATA_FILE)
    ws = wb['Products']
    ws.delete_rows(2, ws.max_row)
    for p in products:
        ws.append([p['id'], p['name'], p['price']])
    wb.save(DATA_FILE)


def read_packs():
    wb = load_workbook(DATA_FILE)
    ws = wb['Packs']
    packs = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            packs.append({'id': row[0], 'name': row[1], 'items': row[2], 'price': row[3]})
    wb.close()
    return packs


def write_packs(packs):
    wb = load_workbook(DATA_FILE)
    ws = wb['Packs']
    ws.delete_rows(2, ws.max_row)
    for pack in packs:
        ws.append([pack['id'], pack['name'], pack['items'], pack['price']])
    wb.save(DATA_FILE)


@app.route('/')
def index():
    return redirect(url_for('list_products'))


@app.route('/products')
def list_products():
    init_workbook()
    products = read_products()
    return render_template('products.html', products=products)


@app.route('/products/new', methods=['GET', 'POST'])
def new_product():
    init_workbook()
    if request.method == 'POST':
        products = read_products()
        next_id = max([p['id'] for p in products], default=0) + 1
        products.append({'id': next_id, 'name': request.form['name'], 'price': request.form['price']})
        write_products(products)
        return redirect(url_for('list_products'))
    return render_template('product_form.html', product=None)


@app.route('/products/<int:product_id>/edit', methods=['GET', 'POST'])
def edit_product(product_id):
    init_workbook()
    products = read_products()
    product = next((p for p in products if p['id'] == product_id), None)
    if not product:
        return redirect(url_for('list_products'))
    if request.method == 'POST':
        product['name'] = request.form['name']
        product['price'] = request.form['price']
        write_products(products)
        return redirect(url_for('list_products'))
    return render_template('product_form.html', product=product)


@app.route('/products/<int:product_id>/delete')
def delete_product(product_id):
    init_workbook()
    products = read_products()
    products = [p for p in products if p['id'] != product_id]
    write_products(products)
    return redirect(url_for('list_products'))


@app.route('/packs')
def list_packs():
    init_workbook()
    packs = read_packs()
    return render_template('packs.html', packs=packs)


@app.route('/packs/new', methods=['GET', 'POST'])
def new_pack():
    init_workbook()
    if request.method == 'POST':
        packs = read_packs()
        next_id = max([p['id'] for p in packs], default=0) + 1
        packs.append({'id': next_id,
                      'name': request.form['name'],
                      'items': request.form['items'],
                      'price': request.form['price']})
        write_packs(packs)
        return redirect(url_for('list_packs'))
    return render_template('pack_form.html', pack=None)


@app.route('/packs/<int:pack_id>/edit', methods=['GET', 'POST'])
def edit_pack(pack_id):
    init_workbook()
    packs = read_packs()
    pack = next((p for p in packs if p['id'] == pack_id), None)
    if not pack:
        return redirect(url_for('list_packs'))
    if request.method == 'POST':
        pack['name'] = request.form['name']
        pack['items'] = request.form['items']
        pack['price'] = request.form['price']
        write_packs(packs)
        return redirect(url_for('list_packs'))
    return render_template('pack_form.html', pack=pack)


@app.route('/packs/<int:pack_id>/delete')
def delete_pack(pack_id):
    init_workbook()
    packs = read_packs()
    packs = [p for p in packs if p['id'] != pack_id]
    write_packs(packs)
    return redirect(url_for('list_packs'))


if __name__ == '__main__':
    init_workbook()
    app.run(debug=True)
